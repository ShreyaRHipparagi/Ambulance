"""Deep inspect of class project ppt - extract ALL style details per slide"""
from pptx import Presentation
from pptx.util import Inches, Pt
import json

pptx_path = "/Users/abhi/Desktop/insta/AMBULANCE/ppt/class project ppt.pptx"
prs = Presentation(pptx_path)

print(f"Slide size: {prs.slide_width.inches:.4f} x {prs.slide_height.inches:.4f} inches")
print(f"Total slides: {len(prs.slides)}\n")

for si, slide in enumerate(prs.slides):
    print(f"\n{'='*70}")
    print(f"SLIDE {si+1}")
    print(f"{'='*70}")
    for sh in slide.shapes:
        left   = round(sh.left/914400,   4)
        top    = round(sh.top/914400,    4)
        width  = round(sh.width/914400,  4)
        height = round(sh.height/914400, 4)
        print(f"\n  [{sh.name}]  type={sh.shape_type}  pos=({left}\", {top}\")  size=({width}\" x {height}\")")
        
        # Fill
        try:
            f = sh.fill
            if f.type == 1:  # SOLID
                print(f"    fill=#{f.fore_color.rgb}")
            elif f.type == 5:
                print(f"    fill=TRANSPARENT")
        except: pass
        
        # Line
        try:
            l = sh.line
            if l.color and l.color.type:
                print(f"    line=#{l.color.rgb}  width={l.width}")
        except: pass

        # Text runs
        if sh.has_text_frame:
            for pi, para in enumerate(sh.text_frame.paragraphs):
                full = para.text.strip()
                if not full: continue
                print(f"    para[{pi}]: align={para.alignment}")
                for ri, run in enumerate(para.runs):
                    t = run.text[:80]
                    sz = run.font.size
                    sz_pt = round(sz/12700, 1) if sz else "inherit"
                    bold = run.font.bold
                    italic = run.font.italic
                    try:    fc = f"#{run.font.color.rgb}"
                    except: fc = "inherit"
                    try:    fn = run.font.name
                    except: fn = "inherit"
                    print(f"      run[{ri}]: '{t}' | sz={sz_pt}pt bold={bold} italic={italic} color={fc} font={fn}")
        
        # Picture
        if sh.shape_type == 13:
            print(f"    [PICTURE SHAPE]")
