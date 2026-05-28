"""
Add an Architecture Diagram slide to the existing PPT,
matching its exact style: dark navy bg, Calibri font, same color palette.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from copy import deepcopy
import copy

# ── Paths ──────────────────────────────────────────────────────────
PPT_IN  = "/Users/abhi/Desktop/insta/AMBULANCE/ppt/class project ppt.pptx"
PPT_OUT = "/Users/abhi/Desktop/insta/AMBULANCE/ppt/class project ppt.pptx"
ARCH_IMG = "/Users/abhi/Desktop/insta/AMBULANCE/architecture_diagram.png"

# ── Exact colors extracted from the existing PPT ───────────────────
BG_NAVY   = RGBColor(0x0B, 0x1D, 0x3A)   # main background
BG_CARD   = RGBColor(0x11, 0x2B, 0x52)   # card / panel fill
CYAN      = RGBColor(0x00, 0xB4, 0xD8)   # primary accent
GREEN     = RGBColor(0x10, 0xB9, 0x81)   # BCS402 green
PURPLE    = RGBColor(0x8B, 0x5C, 0xF6)   # BCS403 purple
AMBER     = RGBColor(0xF5, 0x9E, 0x0B)   # BCS401 amber
WHITE     = RGBColor(0xFF, 0xFF, 0xFF)
TEXT_SOFT = RGBColor(0xE0, 0xF0, 0xFF)   # soft white text
MUTED     = RGBColor(0x64, 0x74, 0x8B)   # footer muted
DARK_BG2  = RGBColor(0x19, 0x30, 0x50)   # footer / bar bg

# ── Load PPT ───────────────────────────────────────────────────────
prs = Presentation(PPT_IN)
W = prs.slide_width.inches    # 10.0"
H = prs.slide_height.inches   # 5.625"

BLANK = prs.slide_layouts[0]  # only one layout in this PPT (DEFAULT)

# ── Helper: add a blank slide ──────────────────────────────────────
def new_slide():
    return prs.slides.add_slide(BLANK)

# ── Helper: rectangle ─────────────────────────────────────────────
def rect(slide, l, t, w, h, fill, border=None, bw=Pt(0), shape_type=1):
    s = slide.shapes.add_shape(shape_type,
                                Inches(l), Inches(t), Inches(w), Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb = fill
    if border:
        s.line.color.rgb = border; s.line.width = bw
    else:
        s.line.fill.background()
    return s

# ── Helper: text box ──────────────────────────────────────────────
def txt(slide, text, l, t, w, h,
        size=12, bold=False, italic=False,
        color=WHITE, align=PP_ALIGN.LEFT, font="Calibri"):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.name  = font
    r.font.size  = Pt(size)
    r.font.bold  = bold
    r.font.italic = italic
    r.font.color.rgb = color

# ── Helper: thin horizontal rule (used in existing slides) ─────────
def hline(slide, l, t, w, color=CYAN):
    s = slide.shapes.add_shape(1, Inches(l), Inches(t), Inches(w), Inches(0))
    s.fill.background()
    s.line.color.rgb = color; s.line.width = Pt(1.5)

# ─────────────────────────────────────────────────────────────────
# BUILD THE ARCHITECTURE SLIDE
# Inserted AFTER slide 1 (index 0) so it sits right after the title
# ─────────────────────────────────────────────────────────────────

# We add it at the end first, then move it to position 1
slide = new_slide()

# 1. Background — exact match to all other slides
rect(slide, 0, 0, W, H, fill=BG_NAVY)

# 2. Decorative top-right corner card (same as slide 1, 8 etc.)
rect(slide, 6.5, -0.5, 4.2, 4.2, fill=BG_CARD, border=CYAN, bw=Pt(1.2))

# 3. Left accent stripe (matches every slide that has a heading)
rect(slide, 0, 0, 0.06, H, fill=CYAN)

# 4. Heading
txt(slide, "System Architecture",
    0.3, 0.15, 7.0, 0.55, size=24, bold=True, color=CYAN)

# 5. Sub-heading line below heading
txt(slide, "End-to-end component interaction & live traffic data flow",
    0.3, 0.7, 8.0, 0.32, size=11, italic=True, color=TEXT_SOFT)

# 6. Thin cyan rule under sub-heading (exact style from other slides)
hline(slide, 0.3, 1.06, 9.4, CYAN)

# 7. Architecture diagram image — left column, large
slide.shapes.add_picture(ARCH_IMG,
                          Inches(0.2), Inches(1.15),
                          Inches(5.9), Inches(4.3))

# 8. Right panel background card
rect(slide, 6.25, 1.15, 3.55, 4.3, fill=BG_CARD, border=CYAN, bw=Pt(0.8))

# 9. Right panel heading
rect(slide, 6.25, 1.15, 3.55, 0.42, fill=CYAN)   # cyan header bar
txt(slide, "Component Legend", 6.35, 1.18, 3.35, 0.34,
    size=11, bold=True, color=BG_NAVY)

# 10. Legend items — match existing color coding
legend = [
    (PURPLE,             "UI Pages  (React JSX)"),
    (CYAN,               "AppContext  — Global State"),
    (CYAN,               "Algorithm Engines  (Dijkstra / A*)"),
    (RGBColor(0xEC,0x48,0x99), "CityGraph  — Adjacency List"),
    (AMBER,              "Mapbox  driving-traffic  API"),
    (GREEN,              "localStorage  Persistence"),
]
for i, (col, label) in enumerate(legend):
    ty = 1.72 + i * 0.52
    # color dot
    s = slide.shapes.add_shape(9,  # oval
                                Inches(6.4), Inches(ty + 0.08),
                                Inches(0.2), Inches(0.2))
    s.fill.solid(); s.fill.fore_color.rgb = col; s.line.fill.background()
    txt(slide, label, 6.7, ty, 3.0, 0.38, size=10, color=TEXT_SOFT)

# 11. Data Flow section inside right panel
ty_flow = 1.72 + len(legend) * 0.52 + 0.1
hline(slide, 6.35, ty_flow, 3.3, CYAN)
txt(slide, "Data Flow", 6.35, ty_flow + 0.08, 3.35, 0.28,
    size=10, bold=True, color=CYAN)
flow_text = (
    "Form → AppContext → "
    "RecommendationService → "
    "Dijkstra SSSP → "
    "MapView → Mapbox API → "
    "Ambulance Animation"
)
txt(slide, flow_text, 6.35, ty_flow + 0.38, 3.35, 1.0,
    size=9, italic=True, color=TEXT_SOFT)

# 12. Bottom badge bar — same as all other slides
rect(slide, 0.3, 5.15, 9.4, 0.3, fill=DARK_BG2)
txt(slide,
    "BCS401 ADA  ·  BCS402 OOCJ  ·  BCS403 DBMS  ·  BCS405B Graph Theory",
    0.3, 5.15, 9.4, 0.3,
    size=10, bold=True, color=CYAN, align=PP_ALIGN.CENTER)

# ─────────────────────────────────────────────────────────────────
# Move the new slide from last position → position index 1
# (right after the title slide, before all existing content slides)
# ─────────────────────────────────────────────────────────────────
from pptx.oxml.ns import qn
from lxml import etree

xml_slides = prs.slides._sldIdLst
# The new slide is the last entry
new_slide_elem = xml_slides[-1]
# Remove from end and insert at index 1
xml_slides.remove(new_slide_elem)
xml_slides.insert(1, new_slide_elem)

# ─────────────────────────────────────────────────────────────────
# Save
# ─────────────────────────────────────────────────────────────────
prs.save(PPT_OUT)
print(f"✅ Saved: {PPT_OUT}")
print(f"   Total slides now: {len(prs.slides)}")
print(f"   Architecture slide inserted at position 2 (index 1)")
