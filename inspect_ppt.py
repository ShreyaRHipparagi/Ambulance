"""Read existing PPT and print all slide details, colors, fonts, sizes"""
from pptx import Presentation
from pptx.util import Inches, Pt
import sys

pptx_path = "/Users/abhi/Desktop/insta/AMBULANCE/ppt/class project ppt.pptx"
prs = Presentation(pptx_path)

print(f"Slide width:  {prs.slide_width.inches:.4f} in")
print(f"Slide height: {prs.slide_height.inches:.4f} in")
print(f"Total slides: {len(prs.slides)}")

for si, slide in enumerate(prs.slides):
    print(f"\n{'='*60}")
    print(f"SLIDE {si+1}  (layout: {slide.slide_layout.name})")
    print(f"{'='*60}")
    for shape in slide.shapes:
        print(f"  Shape: '{shape.shape_type}'  name='{shape.name}'  "
              f"left={shape.left/914400:.3f}\" top={shape.top/914400:.3f}\" "
              f"w={shape.width/914400:.3f}\" h={shape.height/914400:.3f}\"")
        # Fill info
        try:
            fill = shape.fill
            if fill.type is not None:
                print(f"    fill_type={fill.type}")
                try:
                    c = fill.fore_color.rgb
                    print(f"    fill_color=#{c}")
                except: pass
        except: pass
        # Line
        try:
            lc = shape.line.color.rgb
            print(f"    line_color=#{lc}  line_w={shape.line.width}")
        except: pass
        # Text
        if shape.has_text_frame:
            tf = shape.text_frame
            for pi, para in enumerate(tf.paragraphs):
                for ri, run in enumerate(para.runs):
                    print(f"    para[{pi}] run[{ri}]: '{run.text[:60]}' "
                          f"size={run.font.size} bold={run.font.bold} italic={run.font.italic}")
                    try:
                        print(f"      font_color=#{run.font.color.rgb}")
                    except: pass
                    try:
                        print(f"      font_name={run.font.name}")
                    except: pass
        # Picture
        if shape.shape_type == 13:
            print(f"    [PICTURE]")
